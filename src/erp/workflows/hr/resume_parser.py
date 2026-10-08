"""Bounded text extraction for uploaded PDF and DOCX resumes."""

import io
from zipfile import ZipFile

from docx import Document
from pypdf import PdfReader

MAX_RESUME_BYTES = 10 * 1024 * 1024
MAX_RESUME_TEXT_CHARS = 50_000
MAX_PDF_PAGES = 100
MAX_DOCX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_DOCX_ENTRIES = 2_000


def extract_resume_text(filename: str, content_type: str, data: bytes) -> str:
    """Extract text without persisting the uploaded binary."""
    if not data:
        raise ValueError("The uploaded resume is empty.")
    if len(data) > MAX_RESUME_BYTES:
        raise ValueError("Resume files must be 10 MB or smaller.")

    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if suffix == "pdf" and data.startswith(b"%PDF-"):
        if content_type not in {"application/pdf", "application/octet-stream"}:
            raise ValueError("The PDF file content type is not supported.")
        try:
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise ValueError("Encrypted PDF resumes are not supported.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise ValueError(f"PDF resumes may contain at most {MAX_PDF_PAGES} pages.")
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("The PDF resume could not be read.") from exc
    elif suffix == "docx" and data.startswith(b"PK"):
        if content_type not in {
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "application/octet-stream",
        }:
            raise ValueError("The DOCX file content type is not supported.")
        try:
            with ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_DOCX_ENTRIES:
                    raise ValueError("The DOCX resume contains too many embedded files.")
                if sum(entry.file_size for entry in entries) > MAX_DOCX_EXPANDED_BYTES:
                    raise ValueError("The DOCX resume expands beyond the supported size.")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("The DOCX resume is missing its document content.")
            document = Document(io.BytesIO(data))
            blocks = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                blocks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
            text = "\n".join(blocks)
        except ValueError:
            raise
        except Exception as exc:
            raise ValueError("The DOCX resume could not be read.") from exc
    else:
        raise ValueError("Only valid PDF and DOCX resume files are supported.")

    cleaned_text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    cleaned_text = cleaned_text[:MAX_RESUME_TEXT_CHARS]
    if len(cleaned_text.strip()) < 20:
        raise ValueError("No usable text could be extracted from the resume.")
    return cleaned_text
