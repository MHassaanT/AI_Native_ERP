"""Bounded extraction, crawling, and Airtable import helpers for support knowledge."""

from __future__ import annotations

import asyncio
import base64
import email.message
import hashlib
import hmac
import html
import http.client
import io
import ipaddress
import json
import secrets
import socket
import ssl
from datetime import UTC, datetime, timedelta
from html.parser import HTMLParser
from urllib.parse import urljoin, urlsplit
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

import httpx
import xlrd
from cryptography.hazmat.primitives.ciphers.aead import AESGCM
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from lxml.etree import XMLSyntaxError
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PdfReadError
from sqlalchemy.ext.asyncio import AsyncSession
from xlrd.biffh import XLRDError

from erp.config import settings
from erp.db.models.whatsapp import WhatsAppAirtableConnection

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_CHUNKS = 100
CHUNK_SIZE = 5_000
CHUNK_OVERLAP = 300
MAX_EXTRACTED_CHARS = CHUNK_SIZE + (MAX_CHUNKS - 1) * (
    CHUNK_SIZE - 2 * CHUNK_OVERLAP
)
MAX_AIRTABLE_RECORDS = 500
MAX_PDF_PAGES = 200
MAX_DOCX_EXPANDED_BYTES = 100 * 1024 * 1024
MAX_DOCX_ENTRIES = 4_000
MAX_WORKBOOK_CELLS = 250_000
MAX_WEB_RESPONSE_BYTES = 5 * 1024 * 1024
MAX_WEB_REDIRECTS = 5
MAX_AIRTABLE_RESPONSE_BYTES = 5 * 1024 * 1024


class KnowledgeImportError(ValueError):
    """A source could not be safely or usefully imported."""


def _encryption_key() -> bytes:
    key = settings.WHATSAPP_KNOWLEDGE_ENCRYPTION_KEY
    if not key or len(key) != 64:
        raise KnowledgeImportError(
            "Configure WHATSAPP_KNOWLEDGE_ENCRYPTION_KEY as 64 hexadecimal characters."
        )
    try:
        return bytes.fromhex(key)
    except ValueError as exc:
        raise KnowledgeImportError(
            "WHATSAPP_KNOWLEDGE_ENCRYPTION_KEY must contain only hexadecimal characters."
        ) from exc


def encrypt_secret(value: dict) -> str:
    nonce = secrets.token_bytes(12)
    ciphertext = AESGCM(_encryption_key()).encrypt(
        nonce, json.dumps(value, separators=(",", ":")).encode(), None
    )
    return base64.urlsafe_b64encode(nonce + ciphertext).decode()


def decrypt_secret(value: str) -> dict:
    try:
        data = base64.urlsafe_b64decode(value.encode())
        plaintext = AESGCM(_encryption_key()).decrypt(data[:12], data[12:], None)
        result = json.loads(plaintext)
    except KnowledgeImportError:
        raise
    except Exception as exc:
        raise KnowledgeImportError("Stored Airtable credentials could not be decrypted.") from exc
    if not isinstance(result, dict):
        raise KnowledgeImportError("Stored Airtable credentials have an invalid format.")
    return result


def airtable_token_expiry(tokens: dict) -> datetime:
    try:
        expires_in = int(tokens.get("expires_in", 3600))
    except (TypeError, ValueError) as exc:
        raise KnowledgeImportError("Airtable returned an invalid token expiration.") from exc
    if not 1 <= expires_in <= 31_536_000:
        raise KnowledgeImportError("Airtable returned an invalid token expiration.")
    return datetime.now(UTC) + timedelta(seconds=expires_in)


def make_airtable_oauth_state(tenant_id: str, frontend_origin: str, verifier: str) -> str:
    payload = {
        "tenant_id": tenant_id,
        "frontend_origin": frontend_origin,
        "verifier": verifier,
        "issued_at": int(datetime.now(UTC).timestamp()),
        "nonce": secrets.token_urlsafe(18),
    }
    encoded = base64.urlsafe_b64encode(encrypt_secret(payload).encode()).decode().rstrip("=")
    signature = hmac.new(settings.SECRET_KEY.encode(), encoded.encode(), hashlib.sha256).digest()
    return f"{encoded}.{base64.urlsafe_b64encode(signature).decode().rstrip('=')}"


def read_airtable_oauth_state(state: str, allowed_origins: list[str]) -> dict:
    try:
        encoded_text, signature_text = state.split(".", 1)
        encoded = encoded_text.encode()
        signature = base64.urlsafe_b64decode(signature_text + "=" * (-len(signature_text) % 4))
        expected = hmac.new(settings.SECRET_KEY.encode(), encoded, hashlib.sha256).digest()
        if not hmac.compare_digest(signature, expected):
            raise ValueError("Invalid signature.")
        encrypted = base64.urlsafe_b64decode(
            encoded_text + "=" * (-len(encoded_text) % 4)
        ).decode()
        payload = decrypt_secret(encrypted)
        issued_at = datetime.fromtimestamp(int(payload["issued_at"]), UTC)
        if datetime.now(UTC) - issued_at > timedelta(minutes=10):
            raise ValueError("Expired state.")
        if payload["frontend_origin"] not in allowed_origins:
            raise ValueError("Invalid return origin.")
        if not payload.get("nonce") or len(payload.get("verifier", "")) < 43:
            raise ValueError("Invalid state payload.")
        return payload
    except Exception as exc:
        raise KnowledgeImportError("Airtable authorization expired or could not be verified.") from exc


def extract_support_text(filename: str, data: bytes) -> str:
    """Extract bounded text from PDF, DOCX, XLSX/XLS files without persisting binaries."""
    if not data:
        raise KnowledgeImportError("The selected file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise KnowledgeImportError("Knowledge files must be 20 MB or smaller.")

    suffix = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    try:
        if suffix == "pdf":
            if not data.startswith(b"%PDF-"):
                raise KnowledgeImportError("The file is not a valid PDF.")
            reader = PdfReader(io.BytesIO(data))
            if reader.is_encrypted:
                raise KnowledgeImportError("Password-protected PDFs are not supported.")
            if len(reader.pages) > MAX_PDF_PAGES:
                raise KnowledgeImportError(f"PDFs may contain at most {MAX_PDF_PAGES} pages.")
            pieces = [
                f"[Page {number}]\n{page.extract_text() or ''}"
                for number, page in enumerate(reader.pages, start=1)
            ]
            text = "\n\n".join(pieces)
        elif suffix == "docx":
            if not data.startswith(b"PK"):
                raise KnowledgeImportError("The file is not a valid DOCX document.")
            with ZipFile(io.BytesIO(data)) as archive:
                entries = archive.infolist()
                if len(entries) > MAX_DOCX_ENTRIES or sum(
                    entry.file_size for entry in entries
                ) > MAX_DOCX_EXPANDED_BYTES:
                    raise KnowledgeImportError("The DOCX document exceeds safe extraction limits.")
                if "word/document.xml" not in archive.namelist():
                    raise KnowledgeImportError("The DOCX document has no readable content.")
            document = Document(io.BytesIO(data))
            pieces = [paragraph.text for paragraph in document.paragraphs]
            for table in document.tables:
                pieces.extend(
                    " | ".join(cell.text for cell in row.cells) for row in table.rows
                )
            text = "\n".join(pieces)
        elif suffix == "xlsx":
            text = _extract_xlsx(data)
        elif suffix == "xls":
            text = _extract_xls(data)
        else:
            raise KnowledgeImportError("Supported knowledge files are PDF, DOCX, XLSX, and XLS.")
    except KnowledgeImportError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        OSError,
        PackageNotFoundError,
        PdfReadError,
        ValueError,
        XLRDError,
        KeyError,
        IndexError,
        TypeError,
        ParseError,
        XMLSyntaxError,
    ) as exc:
        raise KnowledgeImportError(f"Could not read the {suffix.upper()} file.") from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) > MAX_EXTRACTED_CHARS:
        raise KnowledgeImportError("The file contains more text than can be imported at once.")
    if len(text) < 20:
        raise KnowledgeImportError("No useful text could be extracted from the selected file.")
    return text


def _extract_xlsx(data: bytes) -> str:
    with ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > MAX_DOCX_ENTRIES or sum(entry.file_size for entry in entries) > MAX_DOCX_EXPANDED_BYTES:
            raise KnowledgeImportError("The Excel workbook exceeds safe extraction limits.")
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return _format_workbook(
            (sheet.title, sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def _extract_xls(data: bytes) -> str:
    workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
    try:
        return _format_workbook(
            (
                sheet.name,
                (
                    tuple(sheet.cell_value(row, column) for column in range(sheet.ncols))
                    for row in range(sheet.nrows)
                ),
            )
            for sheet in workbook.sheets()
        )
    finally:
        workbook.release_resources()


def _format_workbook(sheets) -> str:
    lines: list[str] = []
    cell_count = 0
    char_count = 0
    for title, rows in sheets:
        lines.append(f"## Sheet: {title}")
        for row in rows:
            cell_count += len(row)
            if cell_count > MAX_WORKBOOK_CELLS:
                raise KnowledgeImportError("The Excel workbook contains too many cells to import.")
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                line = " | ".join(values)
                char_count += len(line)
                if char_count > MAX_EXTRACTED_CHARS:
                    raise KnowledgeImportError("The Excel workbook contains too much text to import.")
                lines.append(line)
    return "\n".join(lines)


def chunk_knowledge(text: str) -> list[str]:
    """Split normalized text into bounded retrieval passages with small overlap."""
    normalized = text.strip()
    chunks: list[str] = []
    start = 0
    while start < len(normalized) and len(chunks) < MAX_CHUNKS:
        end = min(start + CHUNK_SIZE, len(normalized))
        if end < len(normalized):
            boundary = normalized.rfind(
                "\n", start + CHUNK_SIZE - CHUNK_OVERLAP, end
            )
            if boundary > start:
                end = boundary
        chunk = normalized[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end >= len(normalized):
            start = end
            break
        start = max(end - CHUNK_OVERLAP, start + 1)
    if start < len(normalized) or len(normalized) > MAX_EXTRACTED_CHARS:
        raise KnowledgeImportError("The source exceeds the supported knowledge size.")
    return chunks


def _resolve_public_web_target(raw_url: str) -> tuple[str, str, int, str]:
    try:
        parsed = urlsplit(raw_url.strip())
        scheme = parsed.scheme.lower()
        raw_host = (parsed.hostname or "").rstrip(".").lower()
        host = raw_host.encode("idna").decode("ascii")
        port = parsed.port or (443 if scheme == "https" else 80)
    except ValueError as exc:
        raise KnowledgeImportError("Enter a public HTTP or HTTPS web link.") from exc
    if (
        scheme not in {"http", "https"}
        or not host
        or parsed.username
        or parsed.password
        or port != (443 if scheme == "https" else 80)
    ):
        raise KnowledgeImportError("Enter a public HTTP or HTTPS web link.")
    if host in {"localhost", "localhost.localdomain"} or (
        "." not in host and ":" not in host
    ):
        raise KnowledgeImportError("Local and single-label hosts cannot be crawled.")
    try:
        try:
            address = ipaddress.ip_address(host)
            resolved = {address}
        except ValueError:
            addresses = socket.getaddrinfo(host, port, type=socket.SOCK_STREAM)
            resolved = {ipaddress.ip_address(item[4][0]) for item in addresses}
    except (OSError, ValueError) as exc:
        raise KnowledgeImportError("The web link host could not be resolved.") from exc
    if not resolved or any(not address.is_global for address in resolved):
        raise KnowledgeImportError("Links to private or reserved network addresses are not allowed.")
    return parsed.geturl(), host, port, str(next(iter(resolved)))


async def validate_public_web_url(raw_url: str) -> str:
    """Validate an HTTP(S) URL and reject private, local, or nonstandard targets."""
    target = await asyncio.to_thread(_resolve_public_web_target, raw_url)
    return target[0]


class _PinnedHTTPConnection(http.client.HTTPConnection):
    def __init__(self, host: str, port: int, address: str, timeout: float) -> None:
        super().__init__(host, port, timeout=timeout)
        self._address = address

    def connect(self) -> None:
        self.sock = socket.create_connection((self._address, self.port), self.timeout)


class _PinnedHTTPSConnection(http.client.HTTPSConnection):
    def __init__(
        self, host: str, port: int, address: str, timeout: float, context: ssl.SSLContext
    ) -> None:
        super().__init__(host, port, timeout=timeout, context=context)
        self._address = address

    def connect(self) -> None:
        raw_socket = socket.create_connection((self._address, self.port), self.timeout)
        self.sock = self._context.wrap_socket(raw_socket, server_hostname=self.host)


def _request_pinned_web_page(url: str) -> tuple[int, dict[str, str], bytes]:
    safe_url, host, port, address = _resolve_public_web_target(url)
    parsed = urlsplit(safe_url)
    connection_type = _PinnedHTTPSConnection if parsed.scheme == "https" else _PinnedHTTPConnection
    kwargs = {"context": ssl.create_default_context()} if parsed.scheme == "https" else {}
    connection = connection_type(host, port, address, timeout=15, **kwargs)
    target = parsed.path or "/"
    if parsed.query:
        target += f"?{parsed.query}"
    try:
        connection.request(
            "GET",
            target,
            headers={
                "Accept": "text/html,application/xhtml+xml,text/plain;q=0.9,*/*;q=0.1",
                "Accept-Encoding": "identity",
                "Connection": "close",
                "Host": f"[{host}]" if ":" in host else host,
                "User-Agent": "AI-Native-ERP-Knowledge-Importer/1.0",
            },
        )
        response = connection.getresponse()
        headers = {key.lower(): value for key, value in response.getheaders()}
        content_length = headers.get("content-length")
        if content_length:
            try:
                if int(content_length) > MAX_WEB_RESPONSE_BYTES:
                    raise KnowledgeImportError("Web pages must be 5 MB or smaller.")
            except ValueError as exc:
                raise KnowledgeImportError(
                    "The web page returned an invalid content length."
                ) from exc
        if headers.get("content-encoding", "identity").lower() not in {"", "identity"}:
            raise KnowledgeImportError("Compressed web responses are not supported.")
        body = response.read(MAX_WEB_RESPONSE_BYTES + 1)
        if len(body) > MAX_WEB_RESPONSE_BYTES:
            raise KnowledgeImportError("Web pages must be 5 MB or smaller.")
        return response.status, headers, body
    except OSError as exc:
        raise KnowledgeImportError("Could not fetch the public web page.") from exc
    finally:
        connection.close()


def _fetch_public_web_page(raw_url: str) -> tuple[str, str]:
    current_url = raw_url
    for redirect_count in range(MAX_WEB_REDIRECTS + 1):
        safe_url, _, _, _ = _resolve_public_web_target(current_url)
        status_code, headers, body = _request_pinned_web_page(safe_url)
        if status_code in {301, 302, 303, 307, 308}:
            location = headers.get("location")
            if not location:
                raise KnowledgeImportError("The web page returned an invalid redirect.")
            if redirect_count == MAX_WEB_REDIRECTS:
                raise KnowledgeImportError("The web page redirected too many times.")
            next_url = urljoin(safe_url, location)
            if urlsplit(safe_url).scheme == "https" and urlsplit(next_url).scheme != "https":
                raise KnowledgeImportError("HTTPS pages cannot redirect to an insecure HTTP link.")
            current_url = next_url
            continue
        if status_code != 200:
            raise KnowledgeImportError(f"The web page returned HTTP {status_code}.")
        content_type = headers.get("content-type", "").split(";", 1)[0].strip().lower()
        if content_type not in {"text/html", "application/xhtml+xml", "text/plain"}:
            raise KnowledgeImportError("Only HTML or plain-text web pages can be imported.")
        content_type_header = email.message.Message()
        content_type_header["content-type"] = headers.get("content-type", "")
        charset = content_type_header.get_content_charset() or "utf-8"
        try:
            document = body.decode(charset, errors="replace")
        except LookupError:
            document = body.decode("utf-8", errors="replace")
        if content_type == "text/plain":
            document = f"<html><body><pre>{html.escape(document)}</pre></body></html>"
        return safe_url, document
    raise KnowledgeImportError("The web page redirected too many times.")


class _SafeHTMLForCrawler(HTMLParser):
    _ALLOWED_TAGS = {
        "article", "blockquote", "br", "code", "dd", "div", "dl", "dt", "em",
        "h1", "h2", "h3", "h4", "h5", "h6", "hr", "i", "li", "main", "ol",
        "p", "pre", "section", "small", "span", "strong", "table", "tbody",
        "td", "th", "thead", "tr", "ul",
    }
    _VOID_TAGS = {"br", "hr"}
    _IGNORED_VOID_TAGS = {
        "area", "base", "col", "embed", "img", "input", "link", "meta", "param",
        "source", "track", "wbr",
    }
    _IGNORED_TAGS = {
        "audio", "button", "embed", "form", "head", "iframe", "img", "input",
        "link", "math", "meta", "noscript", "object", "picture", "script", "select",
        "source", "style", "svg", "template", "textarea", "video",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self.ignored_stack: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if self.ignored_stack:
            if tag not in self._IGNORED_VOID_TAGS and tag not in self._VOID_TAGS:
                self.ignored_stack.append(tag)
            return
        if tag in self._IGNORED_TAGS:
            if tag not in self._IGNORED_VOID_TAGS and tag not in self._VOID_TAGS:
                self.ignored_stack.append(tag)
        elif tag in self._ALLOWED_TAGS:
            self.parts.append(f"<{tag}>")

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if not self.ignored_stack and tag in self._ALLOWED_TAGS:
            self.parts.append(f"<{tag}/>")

    def handle_endtag(self, tag: str) -> None:
        if self.ignored_stack:
            if tag in self.ignored_stack:
                index = len(self.ignored_stack) - 1 - self.ignored_stack[::-1].index(tag)
                del self.ignored_stack[index:]
            return
        if tag in self._ALLOWED_TAGS and tag not in self._VOID_TAGS:
            self.parts.append(f"</{tag}>")

    def handle_data(self, data: str) -> None:
        if not self.ignored_stack and data.strip():
            self.parts.append(html.escape(data))


def _sanitize_html_for_crawler(document: str) -> str:
    parser = _SafeHTMLForCrawler()
    parser.feed(document)
    parser.close()
    return "<html><body>" + "".join(parser.parts) + "</body></html>"


async def crawl_with_crawl4ai(url: str) -> str:
    """Fetch a pinned public page and let Crawl4AI format sanitized content as Markdown."""
    try:
        from crawl4ai import AsyncWebCrawler
    except ImportError as exc:
        raise KnowledgeImportError(
            "Web crawling is unavailable because Crawl4AI is not installed on the backend."
        ) from exc
    try:
        _, document = await asyncio.wait_for(
            asyncio.to_thread(_fetch_public_web_page, url), timeout=45
        )
        safe_document = _sanitize_html_for_crawler(document)
        if len(safe_document) < 20:
            raise KnowledgeImportError("The web page did not contain enough readable text.")
        async with AsyncWebCrawler(verbose=False) as crawler:
            result = await asyncio.wait_for(
                crawler.arun(url=f"raw:{safe_document}"), timeout=90
            )
    except TimeoutError as exc:
        raise KnowledgeImportError("The web page took too long to crawl.") from exc
    except KnowledgeImportError:
        raise
    except Exception as exc:
        raise KnowledgeImportError("Crawl4AI could not fetch this web page.") from exc
    markdown = getattr(result, "markdown", None)
    if not isinstance(markdown, str) or len(markdown.strip()) < 20:
        raise KnowledgeImportError("The web page did not contain enough readable text.")
    if len(markdown) > MAX_EXTRACTED_CHARS:
        raise KnowledgeImportError("The web page contains more text than can be imported at once.")
    return markdown


async def airtable_access_token(db: AsyncSession, connection: WhatsAppAirtableConnection) -> str:
    tokens = decrypt_secret(connection.encrypted_tokens)
    expires_at = connection.token_expires_at
    if expires_at and expires_at > datetime.now(UTC) + timedelta(seconds=60):
        access_token = tokens.get("access_token")
        if access_token:
            return str(access_token)

    refresh_token = tokens.get("refresh_token")
    if not isinstance(refresh_token, str) or not refresh_token:
        raise KnowledgeImportError("Airtable authorization expired; reconnect Airtable.")
    response = await _airtable_token_request(
        {
            "grant_type": "refresh_token",
            "refresh_token": refresh_token,
        }
    )
    if not isinstance(response.get("access_token"), str) or not response["access_token"]:
        raise KnowledgeImportError("Airtable did not return a refreshed access token.")
    if response.get("refresh_token"):
        tokens["refresh_token"] = response["refresh_token"]
    tokens["access_token"] = response["access_token"]
    connection.encrypted_tokens = encrypt_secret(tokens)
    connection.token_expires_at = airtable_token_expiry(response)
    await db.flush()
    return str(tokens["access_token"])


async def _airtable_token_request(data: dict[str, str]) -> dict:
    client_id = settings.AIRTABLE_CLIENT_ID
    client_secret = settings.AIRTABLE_CLIENT_SECRET
    if not client_id or not client_secret:
        raise KnowledgeImportError("Airtable OAuth is not configured on this ERP backend.")
    try:
        async with httpx.AsyncClient(timeout=15) as client:
            response = await client.post(
                "https://airtable.com/oauth2/v1/token",
                data=data,
                auth=(client_id, client_secret),
                headers={"Accept": "application/json"},
            )
        response.raise_for_status()
        result = response.json()
    except (httpx.HTTPError, ValueError) as exc:
        raise KnowledgeImportError("Airtable token exchange or refresh failed.") from exc
    if not isinstance(result, dict):
        raise KnowledgeImportError("Airtable returned an invalid OAuth response.")
    return result


async def exchange_airtable_code(code: str, verifier: str) -> dict:
    return await _airtable_token_request(
        {
            "grant_type": "authorization_code",
            "code": code,
            "redirect_uri": settings.AIRTABLE_REDIRECT_URI or "",
            "code_verifier": verifier,
        }
    )


async def airtable_request(
    access_token: str, path: str, *, params: dict | None = None
) -> dict:
    try:
        async with httpx.AsyncClient(timeout=20) as client:
            async with client.stream(
                "GET",
                f"https://api.airtable.com{path}",
                headers={"Authorization": f"Bearer {access_token}"},
                params=params,
            ) as response:
                response.raise_for_status()
                content_length = response.headers.get("content-length")
                if content_length and int(content_length) > MAX_AIRTABLE_RESPONSE_BYTES:
                    raise KnowledgeImportError("Airtable returned a response that is too large.")
                body = bytearray()
                async for part in response.aiter_bytes():
                    body.extend(part)
                    if len(body) > MAX_AIRTABLE_RESPONSE_BYTES:
                        raise KnowledgeImportError("Airtable returned a response that is too large.")
        result = json.loads(body)
    except httpx.HTTPStatusError as exc:
        raise KnowledgeImportError(
            f"Airtable returned HTTP {exc.response.status_code} for the selected resource."
        ) from exc
    except KnowledgeImportError:
        raise
    except (httpx.HTTPError, ValueError, json.JSONDecodeError) as exc:
        raise KnowledgeImportError("Could not reach Airtable.") from exc
    if not isinstance(result, dict):
        raise KnowledgeImportError("Airtable returned an invalid response.")
    return result


async def fetch_airtable_records(
    access_token: str, base_id: str, table_id: str, field_names: list[str]
) -> list[dict]:
    if not field_names:
        raise KnowledgeImportError("Select at least one Airtable field to import.")
    records: list[dict] = []
    offset = None
    while True:
        params: dict = {"pageSize": 100, "fields[]": field_names}
        if offset:
            params["offset"] = offset
        result = await airtable_request(
            access_token,
            f"/v0/{base_id}/{table_id}",
            params=params,
        )
        page = result.get("records", [])
        if not isinstance(page, list) or any(not isinstance(record, dict) for record in page):
            raise KnowledgeImportError("Airtable returned records in an invalid format.")
        records.extend(page)
        if len(records) > MAX_AIRTABLE_RECORDS:
            raise KnowledgeImportError(
                f"Tables with more than {MAX_AIRTABLE_RECORDS} records must be narrowed before importing."
            )
        offset = result.get("offset")
        if not offset:
            return records


def format_airtable_records(table_name: str, records: list[dict], fields: list[str]) -> str:
    lines = [f"# Airtable table: {table_name}"]
    has_selected_values = False
    for record in records:
        record_fields = record.get("fields")
        if not isinstance(record_fields, dict):
            continue
        selected = {
            field: record_fields.get(field)
            for field in fields
            if record_fields.get(field) is not None
        }
        if selected:
            has_selected_values = True
            lines.append(
                json.dumps(selected, ensure_ascii=False, default=str, separators=(",", ":"))
            )
    text = "\n".join(lines)
    if len(text) > MAX_EXTRACTED_CHARS:
        raise KnowledgeImportError("The selected Airtable records contain too much text to import.")
    if not records:
        raise KnowledgeImportError("The selected Airtable table contains no records to import.")
    if not has_selected_values:
        raise KnowledgeImportError("No useful text could be extracted from the selected Airtable fields.")
    return text
