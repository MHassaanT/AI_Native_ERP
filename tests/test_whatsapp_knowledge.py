import io
from types import SimpleNamespace

import crawl4ai
import pytest
from docx import Document
from openpyxl import Workbook
from reportlab.pdfgen import canvas

from erp.config import settings
from erp.workflows.support import whatsapp_knowledge
from erp.workflows.support.whatsapp_knowledge import (
    KnowledgeImportError,
    _sanitize_html_for_crawler,
    chunk_knowledge,
    decrypt_secret,
    encrypt_secret,
    extract_support_text,
    format_airtable_records,
    make_airtable_oauth_state,
    read_airtable_oauth_state,
    validate_public_web_url,
)


@pytest.fixture
def knowledge_encryption_key(monkeypatch):
    monkeypatch.setattr(
        settings,
        "WHATSAPP_KNOWLEDGE_ENCRYPTION_KEY",
        "0123456789abcdef0123456789abcdef0123456789abcdef0123456789abcdef",
    )


def test_extract_support_text_from_pdf():
    buffer = io.BytesIO()
    document = canvas.Canvas(buffer)
    document.drawString(72, 720, "Customer support policies and product information.")
    document.save()

    text = extract_support_text("policies.pdf", buffer.getvalue())

    assert "Customer support policies" in text


def test_extract_support_text_from_docx():
    document = Document()
    document.add_paragraph("Customer support policies and product information.")
    buffer = io.BytesIO()
    document.save(buffer)

    assert "Customer support policies" in extract_support_text("policies.docx", buffer.getvalue())


def test_extract_support_text_from_xlsx():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Support"
    sheet.append(["Question", "Approved answer"])
    sheet.append(["Opening hours", "Monday to Friday, 9 to 5."])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = extract_support_text("faq.xlsx", buffer.getvalue())

    assert "Sheet: Support" in text
    assert "Monday to Friday" in text


def test_rejects_empty_or_unsupported_files():
    with pytest.raises(KnowledgeImportError, match="empty"):
        extract_support_text("empty.pdf", b"")
    with pytest.raises(KnowledgeImportError, match="Supported knowledge files"):
        extract_support_text("notes.txt", b"Some notes with enough content.")


def test_chunks_are_bounded_and_overlap():
    text = "Section one.\n" + ("Customer support detail. " * 400)

    chunks = chunk_knowledge(text)

    assert len(chunks) > 1
    assert all(len(chunk) <= 5_000 for chunk in chunks)

    assert len(chunk_knowledge("x" * 440_600)) <= 100
    with pytest.raises(KnowledgeImportError, match="supported knowledge size"):
        chunk_knowledge("x" * 440_601)


def test_airtable_records_include_only_selected_fields():
    text = format_airtable_records(
        "Support FAQ",
        [{"fields": {"Question": "Hours?", "Answer": "Weekdays.", "Secret": "hidden"}}],
        ["Question", "Answer"],
    )

    assert "Hours?" in text
    assert "Weekdays." in text
    assert "hidden" not in text


def test_airtable_oauth_state_is_encrypted_and_signed(knowledge_encryption_key):
    verifier = "verifier-that-is-long-enough-for-pkce-" + ("a" * 20)
    state = make_airtable_oauth_state(
        "00000000-0000-0000-0000-000000000001",
        "https://erp.example.com",
        verifier,
    )

    assert verifier not in state
    payload = read_airtable_oauth_state(state, ["https://erp.example.com"])
    assert payload["verifier"] == verifier
    assert payload["tenant_id"] == "00000000-0000-0000-0000-000000000001"
    with pytest.raises(KnowledgeImportError, match="could not be verified"):
        read_airtable_oauth_state(state, ["https://other.example.com"])


def test_credentials_encrypt_and_authenticate(knowledge_encryption_key):
    encrypted = encrypt_secret({"access_token": "sensitive-token"})

    assert "sensitive-token" not in encrypted
    assert decrypt_secret(encrypted)["access_token"] == "sensitive-token"


@pytest.mark.asyncio
async def test_web_crawler_rejects_local_targets():
    with pytest.raises(KnowledgeImportError, match="public HTTP or HTTPS"):
        await validate_public_web_url("file:///etc/passwd")
    with pytest.raises(KnowledgeImportError, match="Local and single-label"):
        await validate_public_web_url("http://localhost/admin")
    with pytest.raises(KnowledgeImportError, match="private or reserved"):
        await validate_public_web_url("http://127.0.0.1/")


def test_web_document_sanitizer_removes_scripts_and_network_resources():
    sanitized = _sanitize_html_for_crawler(
        '<html><head><script>fetch("http://127.0.0.1/")</script></head>'
        '<body><article><h1>Support</h1><p>Our hours are 9 to 5.</p>'
        '<img src="http://127.0.0.1/private"><a href="https://example.com">Details</a>'
        '<iframe src="http://169.254.169.254/"></iframe></article></body></html>'
    )

    assert "<h1>Support</h1>" in sanitized
    assert "Our hours are 9 to 5." in sanitized
    assert "fetch(" not in sanitized
    assert "127.0.0.1" not in sanitized
    assert "169.254.169.254" not in sanitized
    assert "href=" not in sanitized


@pytest.mark.asyncio
async def test_web_url_validation_rejects_private_dns_results(monkeypatch):
    monkeypatch.setattr(
        whatsapp_knowledge.socket,
        "getaddrinfo",
        lambda *args, **kwargs: [
            (2, 1, 6, "", ("10.1.2.3", 80)),
        ],
    )

    with pytest.raises(KnowledgeImportError, match="private or reserved"):
        await validate_public_web_url("http://docs.example.test/")


def test_http_redirect_cannot_target_a_private_address(monkeypatch):
    original_resolver = whatsapp_knowledge._resolve_public_web_target

    def fake_resolver(url):
        if url.startswith("https://127.0.0.1"):
            return original_resolver(url)
        return (url, "docs.example.com", 443, "93.184.216.34")

    monkeypatch.setattr(whatsapp_knowledge, "_resolve_public_web_target", fake_resolver)
    monkeypatch.setattr(
        whatsapp_knowledge,
        "_request_pinned_web_page",
        lambda url: (302, {"location": "https://127.0.0.1/admin"}, b""),
    )

    with pytest.raises(KnowledgeImportError, match="private or reserved"):
        whatsapp_knowledge._fetch_public_web_page("https://docs.example.com/")


@pytest.mark.asyncio
async def test_crawl4ai_receives_only_sanitized_local_html(monkeypatch):
    captured = {}

    class FakeCrawler:
        def __init__(self, *, verbose):
            pass

        async def __aenter__(self):
            return self

        async def __aexit__(self, *args):
            return None

        async def arun(self, *, url):
            captured["url"] = url
            return SimpleNamespace(markdown="# Help\n\nOpening hours: 9 to 5.")

    monkeypatch.setattr(
        whatsapp_knowledge,
        "_fetch_public_web_page",
        lambda url: (
            "https://docs.example.com/help",
            '<article><h1>Help</h1><script>alert("unsafe")</script>'
            '<p>Opening hours: 9 to 5.</p></article>',
        ),
    )
    monkeypatch.setattr(crawl4ai, "AsyncWebCrawler", FakeCrawler)

    markdown = await whatsapp_knowledge.crawl_with_crawl4ai("https://docs.example.com/help")

    assert markdown.startswith("# Help")
    assert captured["url"].startswith("raw:<html><body>")
    assert "docs.example.com" not in captured["url"]
    assert "unsafe" not in captured["url"]


@pytest.mark.asyncio
async def test_airtable_record_pagination_and_record_limit(monkeypatch):
    pages = [
        {"records": [{"fields": {"Answer": "A"}}], "offset": "page-2"},
        {"records": [{"fields": {"Answer": "B"}}]},
    ]
    calls = []

    async def fake_request(access_token, path, *, params=None):
        calls.append(params)
        return pages.pop(0)

    monkeypatch.setattr(whatsapp_knowledge, "airtable_request", fake_request)

    records = await whatsapp_knowledge.fetch_airtable_records(
        "token", "appSupport1", "tblQuestions1", ["Answer"]
    )

    assert len(records) == 2
    assert calls[0] == {"pageSize": 100, "fields[]": ["Answer"]}
    assert calls[1] == {"pageSize": 100, "fields[]": ["Answer"], "offset": "page-2"}

    async def too_many_records(access_token, path, *, params=None):
        return {"records": [{}] * 501}

    monkeypatch.setattr(whatsapp_knowledge, "airtable_request", too_many_records)
    with pytest.raises(KnowledgeImportError, match="more than 500 records"):
        await whatsapp_knowledge.fetch_airtable_records(
            "token", "appSupport1", "tblQuestions1", ["Answer"]
        )
